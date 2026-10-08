"""回填往返測試：拆檔 → 模擬組員回填（含各種出錯）→ 收回 → 逐格驗證。

情境：
  A 正常填完 ①②（含 ③）     B 只填 ①（不完整）       C 沒填
  D 改了鎖定欄（竄改）        E 刪掉一列               F 整份沒交回
  G 一份經 LibreOffice 另存（模擬 Excel / Google 另存後格式變動）
回傳 (ok: bool, 說明: str)。gate.py 的 G8 呼叫本檔。
"""
from __future__ import annotations

import random
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "flow"))
import backfill as bf  # noqa: E402

FIXTURE = ROOT / "tests" / "fixtures" / "T組進度完整分析_0916.xlsx"


def run(verbose=False):
    rnd = random.Random(916)
    tmp = Path(tempfile.mkdtemp(prefix="bf_rt_"))
    try:
        out = tmp / "split"
        import contextlib
        import io
        with contextlib.redirect_stdout(io.StringIO()):
            manifest = bf.split(FIXTURE, out)
        files = sorted(out.glob("*.xlsx"))
        if len(files) < 3:
            return False, f"拆檔只有 {len(files)} 份"
        back = tmp / "back"
        back.mkdir()
        expect_vals = {}            # key -> {col_idx: value}
        expect_status = Counter()
        tamper_keys, lost_keys = set(), set()

        missing_file = files[-1]                     # F：最後一份不交回
        lo_file = files[1]                           # G：這份經 LibreOffice 另存
        C = bf.cfg()
        for f in files:
            wb = openpyxl.load_workbook(f)
            owner_all_keys = []
            for ws in wb.worksheets[1:]:
                hdr = [bf.norm(c.value) for c in ws[1]]
                ki = hdr.index(bf.KEY_COL)
                fill = bf.col_index(hdr[:ki], C["fill_columns"])
                req = bf.required_idx(hdr[:ki], C)
                opt = [i for i in fill if i not in req]
                rows_to_delete = []
                for r in range(2, ws.max_row + 1):
                    key = ws.cell(r, ki + 1).value
                    owner_all_keys.append(key)
                    p = rnd.random()
                    if f == missing_file:
                        continue
                    if p < 0.02 and len(lost_keys) < 3:          # E 刪列
                        rows_to_delete.append(r)
                        lost_keys.add(key)
                        continue
                    if p < 0.04 and len(tamper_keys) < 3:        # D 竄改鎖定欄
                        lock_col = next(j for j in range(ki) if j not in fill)
                        ws.cell(r, lock_col + 1).value = "被改了"
                        ws.cell(r, req[0] + 1).value = "có điền nhưng sửa ô khóa"
                        tamper_keys.add(key)
                        continue
                    if p < 0.75:                                 # A
                        vals = {req[0]: f"Đã kiểm tra {r}", req[-1]: f"KH xác nhận, xử lý {r}"}
                        if opt and p < 0.3:
                            vals[opt[0]] = f"CT-{r:05d}"
                        for i, v in vals.items():
                            ws.cell(r, i + 1).value = v
                        expect_vals[key] = vals
                    elif p < 0.85 and len(req) > 1:              # B 只填 ①
                        ws.cell(r, req[0] + 1).value = f"Đang kiểm tra {r}"
                        expect_vals[key] = {req[0]: f"Đang kiểm tra {r}"}
                    # C：不填
                for r in reversed(rows_to_delete):
                    ws.delete_rows(r)
            if f == missing_file:
                continue
            dst = back / f.name
            wb.save(dst)
            if f == lo_file and shutil.which("soffice"):  # G（沒裝 LibreOffice 的電腦略過此情境）
                lo = tmp / "lo"
                r = subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(lo), str(dst)],
                                   capture_output=True, text=True, timeout=180)
                if r.returncode == 0 and (lo / f.name).exists():
                    shutil.copy(lo / f.name, dst)
                else:
                    return False, f"LibreOffice 另存失敗：{r.stderr[-200:]}"

        res = tmp / "res"
        with contextlib.redirect_stdout(io.StringIO()):
            cnt, issues, merged = bf.collect(FIXTURE, back, res)

        # 驗證 1：每個預期寫回值都在原位
        mb = openpyxl.load_workbook(merged, data_only=True)
        bad = 0
        for key, vals in expect_vals.items():
            _t, title, r = key.split("|")
            for i, v in vals.items():
                if bf.norm(mb[title].cell(int(r), i + 1).value) != v:
                    bad += 1
        if bad:
            return False, f"{bad} 格回填值沒寫到原位"
        # 驗證 2：竄改列一格都沒寫回
        for key in tamper_keys:
            _t, title, r = key.split("|")
            hdr = [bf.norm(c.value) for c in mb[title][bf.find_header(mb[title], C['fill_columns'])[0]]]
            for i in bf.col_index(hdr, C["fill_columns"]):
                orig = openpyxl.load_workbook(FIXTURE, data_only=True)[title].cell(int(r), i + 1).value
                if bf.norm(mb[title].cell(int(r), i + 1).value) != bf.norm(orig):
                    return False, f"竄改列 {key} 被寫回了"
        # 驗證 3：狀態統計
        kinds = Counter(i[0] for i in issues)
        if kinds["鎖定欄被改"] != len(tamper_keys):
            return False, f"竄改偵測 {kinds['鎖定欄被改']} ≠ 實際 {len(tamper_keys)}"
        if cnt["列遺失"] != len(lost_keys):
            return False, f"列遺失 {cnt['列遺失']} ≠ 實際 {len(lost_keys)}"
        if cnt["未交回"] == 0:
            return False, "整份未交回沒有被抓到"
        if cnt["已回填"] == 0 or cnt["不完整"] == 0:
            return False, f"狀態分布異常：{dict(cnt)}"
        # 驗證 4（collect 內建）：回填欄以外逐格不變；能走到這裡代表已通過
        msg = (f"{len(files)} 份拆檔｜寫回 {sum(len(v) for v in expect_vals.values())} 格全對｜"
               f"竄改 {len(tamper_keys)}、刪列 {len(lost_keys)}、未交回 1 份"
               + ("、LibreOffice 另存 1 份" if shutil.which("soffice") else "（本機無 LibreOffice，另存情境略過）")
               + " 皆正確處理")
        return True, msg
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    ok, msg = run(verbose=True)
    print(("✅ " if ok else "❌ ") + msg)
    sys.exit(0 if ok else 1)
