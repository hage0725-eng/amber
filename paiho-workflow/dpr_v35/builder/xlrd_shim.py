# FILE_VERSION = 1   # R-0911-7（Amber 林彥博 2026-09-11「A5-a」·轉正·永久）：
#                    xlrd 相容墊片收入 _PACK 第 9 支核心檔。改本檔須同步遞增此常數
#                    （R-v555-1 降版閘門依此取版號；帶標記故不屬 R-0827-4 高風險檔）。
# -*- coding: utf-8 -*-
"""xlrd 相容墊片（shim） — 本容器 pypi 出口被封鎖，xlrd 無法安裝。
以 LibreOffice(.xls→.xlsx) + openpyxl 重現 build_dpr.py 所用之 xlrd 極小 API：
  open_workbook(path) / sheet_by_name / sheet_by_index / nrows / ncols / cell_value(r,c)
語意對齊 xlrd：日期/時間儲存格一律回傳 Excel 序列值 float；空白儲存格回傳 ''。
"""
import datetime as _dt
import hashlib as _hl
import os as _os
import subprocess as _sp

import openpyxl as _op

_EPOCH = _dt.datetime(1899, 12, 30)
_CACHE = '/home/claude/dpr/_xlsconv'


def _to_serial(v):
    if isinstance(v, _dt.datetime):
        return (v - _EPOCH).total_seconds() / 86400.0
    if isinstance(v, _dt.date):
        return float((_dt.datetime(v.year, v.month, v.day) - _EPOCH).days)
    if isinstance(v, _dt.time):
        return (v.hour * 3600 + v.minute * 60 + v.second) / 86400.0
    if isinstance(v, _dt.timedelta):
        return v.total_seconds() / 86400.0
    return v


def _convert(path):
    _os.makedirs(_CACHE, exist_ok=True)
    key = _hl.md5((_os.path.abspath(path) + str(_os.path.getmtime(path))).encode()).hexdigest()[:12]
    out = _os.path.join(_CACHE, key + '.xlsx')
    if not _os.path.exists(out):
        tmp = _os.path.join(_CACHE, key)
        _os.makedirs(tmp, exist_ok=True)
        r = _sp.run(['libreoffice', '--headless', '--convert-to', 'xlsx',
                     '--outdir', tmp, path], capture_output=True, timeout=900)
        cands = [f for f in _os.listdir(tmp) if f.lower().endswith('.xlsx')]
        if not cands:
            raise RuntimeError('LibreOffice 轉檔失敗: %s / %s' % (path, r.stderr[-300:]))
        _os.replace(_os.path.join(tmp, cands[0]), out)
    return out


class Sheet:
    def __init__(self, ws):
        self._ws = ws
        rows = [[_to_serial(c) for c in row] for row in ws.iter_rows(values_only=True)]
        while rows and all(v is None or v == '' for v in rows[-1]):
            rows.pop()
        ncols = 0
        for row in rows:
            for i in range(len(row) - 1, -1, -1):
                if row[i] is not None and row[i] != '':
                    ncols = max(ncols, i + 1)
                    break
        self._rows = [list(r) + [''] * (ncols - len(r)) for r in rows]
        for r in self._rows:
            for i, v in enumerate(r):
                if v is None:
                    r[i] = ''
        self.nrows = len(self._rows)
        self.ncols = ncols
        self.name = ws.title

    def cell_value(self, r, c):
        return self._rows[r][c]

    def row_values(self, r):
        return list(self._rows[r])


class Book:
    def __init__(self, path):
        p = path if path.lower().endswith(('.xlsx', '.xlsm')) else _convert(path)
        self._wb = _op.load_workbook(p, data_only=True, read_only=False)
        self._cache = {}

    def sheet_names(self):
        return list(self._wb.sheetnames)

    def sheet_by_name(self, name):
        if name not in self._cache:
            self._cache[name] = Sheet(self._wb[name])
        return self._cache[name]

    def sheet_by_index(self, i):
        return self.sheet_by_name(self._wb.sheetnames[i])


def open_workbook(path, *a, **kw):
    return Book(path)
