"""回填分派與收回合併。

split  ：把 DPR 產出的分析檔，按「助理（無則業務）」拆成每人一份回填檔。
collect：把交回的回填檔合併回分析檔原位（格式不變，DPR 照常讀），並產出回填狀態表。

用法：
  python flow/backfill.py split   <分析檔.xlsx> <輸出資料夾>
  python flow/backfill.py collect <分析檔.xlsx> <交回資料夾> <輸出資料夾>

設計原則：
  * 每列帶隱藏欄 _KEY（分頁|原列號）與 _SIG（鎖定欄位簽章），收回時以此定位與防竄改。
  * 只寫回「回填欄」，其他儲存格一格不動；合併後逐格比對原檔驗證。
  * 不做語意判斷（原因分類、駁回）——那是 DPR 字典與 classifier 的工作，避免兩套規則。
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import os
import re
import sys
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "core"))
from paiho_core import load_rules  # noqa: E402

BLACK = Side(style="thin", color="000000")
BORDER = Border(left=BLACK, right=BLACK, top=BLACK, bottom=BLACK)
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="1F4E78")
KEY_COL, SIG_COL, ORIG_COL = "_KEY", "_SIG", "_ORIG"


# ---------------------------------------------------------------- 共用
def cfg():
    return load_rules()["flow"]["backfill"]


def slug(name: str) -> str:
    """越南名 → 純英數檔名（下載連結不亂碼）。"""
    s = unicodedata.normalize("NFKD", name.replace("Đ", "D").replace("đ", "d"))
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-")
    return s or "unassigned"


def norm(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (dt.datetime, dt.date)):
        v = v.isoformat()[:10]
    return str(v).strip()


def find_header(ws, keys, max_scan=6):
    for r in range(1, min(max_scan, ws.max_row) + 1):
        vals = [norm(c.value) for c in ws[r]]
        if any(any(k in v for k in keys) for v in vals):
            return r, vals
    return None, None


def backfill_sheets(wb, C):
    for ws in wb.worksheets:
        if any(ws.title.startswith(p) for p in C["sheets"]):
            hr, hdr = find_header(ws, C["fill_columns"])
            if hr:
                yield ws, hr, hdr


def col_index(hdr, keys):
    return [i for i, v in enumerate(hdr) if any(v.startswith(k) for k in keys)]


def required_idx(hdr, C):
    """必填欄：有①②就用①②；像 8.NTNS 沒有①②的分頁，改用第一個回填欄（NT/NS判定）。"""
    req = col_index(hdr, C["required_columns"])
    return req or col_index(hdr, C["fill_columns"])[:1]


BAD_OWNER = {"", "0", "None", "nan", "—", "-", "#N/A", "#REF!", "#VALUE!", "#NAME?", "#DIV/0!", "#NULL!", "#NUM!"}


def owner_of(row_vals, hdr, C):
    for k in C["owner_columns"]:
        idx = col_index(hdr, [k])
        if idx:
            v = unicodedata.normalize("NFC", norm(row_vals[idx[0]]))
            if v not in BAD_OWNER and not v.startswith("#"):
                return v
    return C["unassigned"]


def signature(row_vals, fill_idx) -> str:
    locked = [norm(v) for i, v in enumerate(row_vals) if i not in fill_idx]
    return hashlib.sha1("\x1f".join(locked).encode("utf-8")).hexdigest()[:16]


def is_data_row(row_vals):
    vals = [norm(v) for v in row_vals]
    filled = [v for v in vals if v]
    if not filled:
        return False
    # 佔位列：「本輪無…」「（佔位列…」這類只有說明文字
    return not (len(filled) <= 2 and any(t in filled[0] for t in ("本輪", "佔位", "無異常")))


# ---------------------------------------------------------------- split
def split(src, out_dir):
    C = cfg()
    src, out_dir = Path(src), Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(src, data_only=True)
    tag = re.search(r"(\d{4})(?!.*\d{4})", src.stem)
    tag = tag.group(1) if tag else dt.date.today().strftime("%m%d")

    buckets = defaultdict(lambda: defaultdict(list))   # owner -> sheet -> rows
    headers = {}
    for ws, hr, hdr in backfill_sheets(wb, C):
        fill_idx = set(col_index(hdr, C["fill_columns"]))
        headers[ws.title] = (hdr, sorted(fill_idx))
        for r in range(hr + 1, ws.max_row + 1):
            vals = [c.value for c in ws[r]][: len(hdr)]
            if not is_data_row(vals):
                continue
            key = f"{tag}|{ws.title}|{r}"
            orig = json.dumps({str(i): norm(vals[i]) for i in sorted(fill_idx) if i < len(vals)}, ensure_ascii=False)
            buckets[owner_of(vals, hdr, C)][ws.title].append((key, signature(vals, fill_idx), orig, vals))

    manifest, used = [], {}
    for owner, sheets in sorted(buckets.items()):
        base = f"DPR_{tag}_{slug(owner)}"
        n = used.get(base, 0) + 1
        used[base] = n
        name = f"{base}.xlsx" if n == 1 else f"{base}-{n}.xlsx"     # 不同人轉成同一檔名時加序號，絕不覆蓋
        if n > 1:
            print(f"⚠️ 檔名撞名：{owner} → {name}")
        ob = openpyxl.Workbook()
        _readme(ob.active, owner, tag, sum(len(v) for v in sheets.values()))
        for title, rows in sheets.items():
            hdr, fill_idx = headers[title]
            ws = ob.create_sheet(title[:31])
            ws.sheet_view.showGridLines = True
            ws.freeze_panes = None
            full = hdr + [KEY_COL, SIG_COL, ORIG_COL]
            for j, h in enumerate(full, 1):
                c = ws.cell(1, j, h)
                c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), HEAD, BORDER
                c.alignment = Alignment(wrap_text=True, vertical="center")
            for i, (key, sig, orig, vals) in enumerate(rows, 2):
                for j, v in enumerate(list(vals) + [key, sig, orig], 1):
                    c = ws.cell(i, j, v)
                    c.border = BORDER
                    if (j - 1) in fill_idx:
                        c.fill = YELLOW
                        c.protection = Protection(locked=False)
                        c.alignment = Alignment(wrap_text=True, vertical="top")
            for j, h in enumerate(full, 1):
                w = 40 if (j - 1) in fill_idx else min(max(len(norm(h)) * 0.9, 10), 28)
                ws.column_dimensions[get_column_letter(j)].width = w
            for j in (len(hdr) + 1, len(hdr) + 2, len(hdr) + 3):
                ws.column_dimensions[get_column_letter(j)].hidden = True
            ws.auto_filter.ref = f"A1:{get_column_letter(len(full))}{len(rows) + 1}"
            ws.protection.sheet = True
            ws.protection.password = C["protect_password"]
            ws.protection.autoFilter = False      # False＝允許篩選
            ws.protection.sort = True
            ws.protection.formatColumns = False
            ws.protection.formatRows = False
        ob.save(out_dir / name)
        manifest.append((owner, name, {k: len(v) for k, v in sheets.items()}))

    print(f"拆檔完成：{len(manifest)} 份 → {out_dir}")
    for owner, name, cnt in manifest:
        print(f"  {name}  {owner}  " + "、".join(f"{k.split()[0]}={v}" for k, v in cnt.items()))
    return manifest


def _readme(ws, owner, tag, n):
    ws.title = "說明 Hướng dẫn"
    ws.sheet_view.showGridLines = True
    lines = [
        (f"T組回填檔 {tag[:2]}/{tag[2:]}｜{owner}", f"Tệp phản hồi nhóm T ngày {tag[2:]}/{tag[:2]} | {owner}"),
        (f"共 {n} 列待回填", f"Tổng cộng {n} dòng cần phản hồi"),
        ("只需填寫黃底欄位：①異常原因檢查 ②處理方式和結果（必填）③憑證單號（選填）",
         "Chỉ điền các cột nền vàng: ① Báo cáo kiểm tra nguyên nhân ② Cách xử lý và kết quả (bắt buộc) ③ Mã chứng từ (không bắt buộc)"),
        ("其他欄位已鎖定，請勿刪列、插列或修改；可使用篩選",
         "Các cột khác đã được khóa; vui lòng không xóa dòng, chèn dòng hoặc chỉnh sửa. Có thể dùng bộ lọc."),
        ("填完後存檔，保持原檔名放回原資料夾即可",
         "Sau khi điền xong, vui lòng lưu tệp và giữ nguyên tên tệp trong thư mục ban đầu."),
    ]
    for i, (zh, vi) in enumerate(lines, 1):
        ws.cell(i, 1, zh).border = BORDER
        ws.cell(i, 2, vi).border = BORDER
    ws.column_dimensions["A"].width = 70
    ws.column_dimensions["B"].width = 90


# ---------------------------------------------------------------- collect
def master_tag(master: Path) -> str:
    m = re.search(r"(\d{4})(?!.*\d{4})", Path(master).stem)
    return m.group(1) if m else ""


def collect(master, returned_dir, out_dir, expected_owners=None):
    C = cfg()
    master, returned_dir, out_dir = Path(master), Path(returned_dir), Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    tag = master_tag(master)
    wb = openpyxl.load_workbook(master)                 # 保留格式
    wbv = openpyxl.load_workbook(master, data_only=True)

    sheet_meta, expected = {}, {}                       # key -> (owner, sig, vals)
    for ws, hr, hdr in backfill_sheets(wbv, C):
        fill_idx = col_index(hdr, C["fill_columns"])
        sheet_meta[ws.title] = (hdr, fill_idx, required_idx(hdr, C))
        for r in range(hr + 1, ws.max_row + 1):
            vals = [c.value for c in ws[r]][: len(hdr)]
            if is_data_row(vals):
                expected[f"{tag}|{ws.title}|{r}"] = (owner_of(vals, hdr, C), signature(vals, set(fill_idx)), vals)

    issues = []
    def issue(kind, key, fname, note):
        owner, ref = "", key
        if key in expected:
            owner, ref = expected[key][0], _order_ref(expected[key][2], sheet_meta[key.split("|")[1]][0])
        issues.append((kind, key, fname, note, owner, ref))

    # 交回檔：只收 .xlsx；其他格式一律列為異常，不默默略過。同一列多份交回 → 取最新修改的檔
    files = []
    for f in sorted(returned_dir.iterdir()):
        if f.is_dir() or f.name.startswith(("~$", ".")):
            continue
        if f.suffix.lower() != ".xlsx":
            issue("無法讀取的檔案", f.name, f.name, "只收 .xlsx；請另存為 Excel 活頁簿後放回")
            continue
        files.append(f)
    files.sort(key=lambda f: os.path.getmtime(f), reverse=True)

    got, src_of, tampered, owners_back = {}, {}, set(), set()
    for f in files:
        try:
            rb = openpyxl.load_workbook(f, data_only=True)
        except Exception as e:
            issue("無法讀取的檔案", f.name, f.name, f"開檔失敗：{type(e).__name__}")
            continue
        for ws in rb.worksheets:
            hr, hdr = find_header(ws, [KEY_COL])
            if not hr or KEY_COL not in hdr or SIG_COL not in hdr or ORIG_COL not in hdr:
                continue
            ki, si, oi = hdr.index(KEY_COL), hdr.index(SIG_COL), hdr.index(ORIG_COL)
            for r in range(hr + 1, ws.max_row + 1):
                vals = [c.value for c in ws[r]]
                key = norm(vals[ki]) if ki < len(vals) else ""
                if not key:
                    continue
                if not key.startswith(f"{tag}|"):
                    issue("日期不符", key, f.name, f"這是 {key.split('|')[0]} 的回填檔，本次彙總 {tag}，不寫回")
                    continue
                if key not in expected:
                    issue("無對應列", key, f.name, "列號對不上分析檔，不寫回")
                    continue
                owners_back.add(expected[key][0])
                title = key.split("|")[1]
                mhdr, fill_idx, _ = sheet_meta[title]
                body = (vals[: len(mhdr)] + [None] * len(mhdr))[: len(mhdr)]
                if signature(body, set(fill_idx)) != expected[key][1] or norm(vals[si]) != expected[key][1]:
                    issue("鎖定欄被改", key, f.name, "此列不寫回，請本人確認")
                    tampered.add(key)
                    continue
                if key in got:
                    if any(norm(got[key][i]) != norm(body[i]) for i in fill_idx):
                        issue("重複交回內容不同", key, f.name, f"採較新的 {src_of[key]}")
                    continue
                try:
                    orig = {int(k): v for k, v in json.loads(norm(vals[oi])).items()}
                except Exception:
                    issue("鎖定欄被改", key, f.name, "隱藏欄 _ORIG 損毀，此列不寫回")
                    tampered.add(key)
                    continue
                got[key] = {i: body[i] for i in fill_idx}
                got[key]["_orig"] = orig
                src_of[key] = f.name

    # 寫回：只寫「組員有改的格」；Amber 拆檔後自己改過的格，組員又改＝衝突，保留 Amber 的
    planned = {}                                        # (title, r, col) -> 最終值
    status = []
    for key, (owner, _sig, vals) in expected.items():
        _t, title, r = key.split("|")
        r = int(r)
        mhdr, fill_idx, req_idx = sheet_meta[title]
        ws = wb[title]
        if key in got:
            orig = got[key]["_orig"]
            final = {}
            for i in fill_idx:
                ret, o, cur = norm(got[key][i]), norm(orig.get(i, "")), norm(ws.cell(r, i + 1).value)
                if ret == o:
                    final[i] = cur
                elif cur == o:
                    final[i] = ret
                    if ret:
                        v = got[key][i]
                        if isinstance(v, (dt.datetime, dt.date)):
                            v = norm(v)                 # 日期一律存成文字，避免格式漂移
                        ws.cell(r, i + 1).value = v
                        if isinstance(v, str) and v.startswith("="):
                            ws.cell(r, i + 1).data_type = "s"     # 「==> 已處理」不可變公式
                    else:
                        ws.cell(r, i + 1).value = None            # 組員清空
                elif ret != cur:
                    issue("與 Amber 修改衝突", key, src_of[key], "拆檔後 Amber 已改此格，保留 Amber 版本")
                    final[i] = cur
                else:
                    final[i] = cur
                planned[(title, r, i + 1)] = final[i]
            req_filled = [bool(final.get(i)) for i in req_idx]
            st = "已回填" if all(req_filled) else ("不完整" if any(req_filled) else "未回填")
        elif key in tampered:
            st = "鎖定欄被改"
        elif owner in owners_back:
            st = "列遺失"
        else:
            st = "未交回"
        status.append((owner, title, r, st, _order_ref(vals, mhdr)))

    stem = master.stem
    out_master = out_dir / f"{stem}_回填彙總.xlsx"
    tmp = out_dir / f".{stem}_回填彙總.tmp.xlsx"
    wb.save(tmp)
    try:
        _verify(master, tmp, sheet_meta, planned)
    except SystemExit:
        tmp.unlink(missing_ok=True)
        raise
    os.replace(tmp, out_master)                         # 驗證過才換上正式檔名
    _status_book(out_dir / f"{stem}_回填狀態.xlsx", status, issues)

    from collections import Counter
    cnt = Counter(s[3] for s in status)
    print(f"收回完成：讀入 {len(files)} 份檔")
    print("  " + "、".join(f"{k} {v}" for k, v in cnt.most_common()) + f"、異常 {len(issues)}")
    print(f"  ==> {out_master.name}")
    return cnt, issues, out_master


def _order_ref(vals, hdr):
    def g(k):
        idx = col_index(hdr, [k])
        return norm(vals[idx[0]]) if idx else ""
    order = g("訂單號碼") or g("89訂單號") or g("料號")
    return f"{g('客戶代碼')} {g('客戶簡稱')}｜{order}".strip()


def _verify(src, out, sheet_meta, planned):
    """合併後逐格比對：回填欄以外完全不變；回填欄＝預定的最終值（含寫錯列也抓得到）。"""
    a = openpyxl.load_workbook(src, data_only=True)
    b = openpyxl.load_workbook(out)
    if a.sheetnames != b.sheetnames:
        raise SystemExit("合併後分頁順序改變，中止輸出")
    for ws in a.worksheets:
        wb2 = b[ws.title]
        for row in ws.iter_rows():
            for c in row:
                got = wb2.cell(c.row, c.column)
                want = planned.get((ws.title, c.row, c.column), norm(c.value))
                if norm(got.value) != want:
                    raise SystemExit(f"合併驗證失敗：{ws.title}!{c.coordinate} 應為「{want[:20]}」實為「{norm(got.value)[:20]}」，中止輸出")
                if got.data_type == "f" and not (isinstance(c.value, str) and c.value.startswith("=")):
                    raise SystemExit(f"合併驗證失敗：{ws.title}!{c.coordinate} 變成公式，中止輸出")


def _status_book(path, status, issues):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "回填狀態 Trạng thái"
    from collections import Counter
    by = defaultdict(Counter)
    for owner, _t, _r, st, _ref in status:
        by[owner][st] += 1
    cols = ["已回填", "不完整", "未回填", "未交回", "鎖定欄被改", "列遺失"]
    _table(ws, ["助理 / Trợ lý"] + cols + ["合計 / Tổng"],
           [[o] + [c[k] for k in cols] + [sum(c.values())] for o, c in sorted(by.items())])
    ws2 = wb.create_sheet("待追 Cần theo dõi")
    _table(ws2, ["助理 / Trợ lý", "分頁 / Bảng", "列 / Dòng", "狀態 / Trạng thái", "客代 簡稱｜單號"],
           [list(s) for s in status if s[3] != "已回填"])
    ws3 = wb.create_sheet("異常 Bất thường")
    _table(ws3, ["類型 / Loại", "鍵 / Khóa", "檔案 / Tệp", "說明 / Ghi chú", "助理 / Trợ lý", "客代 簡稱｜單號"],
           [list(i) for i in issues])
    wb.save(path)


def _table(ws, head, rows):
    ws.sheet_view.showGridLines = True
    ws.freeze_panes = None
    for j, h in enumerate(head, 1):
        c = ws.cell(1, j, h)
        c.font, c.fill, c.border = Font(bold=True, color="FFFFFF"), HEAD, BORDER
    for i, row in enumerate(rows, 2):
        for j, v in enumerate(row, 1):
            ws.cell(i, j, v).border = BORDER
    for j in range(1, len(head) + 1):
        ws.column_dimensions[get_column_letter(j)].width = 22 if j > 1 else 30


if __name__ == "__main__":
    a = sys.argv[1:]
    if len(a) == 3 and a[0] == "split":
        split(a[1], a[2])
    elif len(a) == 4 and a[0] == "collect":
        collect(a[1], a[2], a[3])
    else:
        raise SystemExit(__doc__)
